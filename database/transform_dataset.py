import shutil
from pathlib import Path
import openpyxl

BASE_DIR = Path(r"C:\Users\MANIKANDAN G\Documents\WE LAB Project")
XLSX_PATH = BASE_DIR / "frontend" / "data" / "craftora_1000_category_wise_handicraft_dataset.xlsx"
BACKUP_PATH = BASE_DIR / "frontend" / "data" / "craftora_1000_backup.xlsx"

# 1. Make a backup
if not BACKUP_PATH.exists():
    shutil.copy2(XLSX_PATH, BACKUP_PATH)
    print("Backup created at", BACKUP_PATH)

wb = openpyxl.load_workbook(str(XLSX_PATH))
ws = wb.active

# Dictionary of replacements: old_name -> dict of new fields
REPLACEMENTS = {
    # ── MEN CATEGORY ──────────────────────────────────────────────
    "Leather Wallet": {
        "name": "Shantiniketan Embossed Leather Folio",
        "material": "Embossed Leather",
        "technique": "Shantiniketan leathercraft",
        "state": "West Bengal",
        "primary_use": "Personal & office use",
        "desc": "Handcrafted Shantiniketan embossed leather folio featuring traditional floral and tribal motifs, handcrafted by rural artisans of Birbhum using vegetable dyes."
    },
    "Premium Leather Belt": {
        "name": "Hand-embroidered Kashmiri Aari Waistcoat",
        "material": "Raw Silk & Wool",
        "technique": "Aari hand embroidery",
        "state": "Jammu and Kashmir",
        "primary_use": "Festive wear",
        "desc": "Traditional handcrafted Kashmiri men's waistcoat (bandi) featuring intricate floral Aari thread embroidery handcrafted on pure raw silk fabric."
    },
    "Premium Handcrafted Tie": {
        "name": "Handloom Pashmina Wool Shawl for Men",
        "material": "Pure Pashmina Wool",
        "technique": "Handloom weaving",
        "state": "Jammu and Kashmir",
        "primary_use": "Ethnic wear",
        "desc": "Ultra-soft handloom Pashmina wool shawl for men, handwoven on traditional wooden looms in the Kashmir valley with delicate hashia borders."
    },
    "Floral Leather Card Holder": {
        "name": "Rajasthani Handcrafted Camel Leather Mojari",
        "material": "Camel Leather",
        "technique": "Mojari leathercraft",
        "state": "Rajasthan",
        "primary_use": "Ethnic footwear",
        "desc": "Authentic Rajasthani handcrafted camel leather mojari ethnic shoes featuring fine silk thread embroidery and cushioned comfort."
    },
    "Festive Leather Card Holder": {
        "name": "Rajasthani Embroidered Camel Leather Mojari",
        "material": "Camel Leather",
        "technique": "Mojari leathercraft",
        "state": "Rajasthan",
        "primary_use": "Ethnic footwear",
        "desc": "Festive hand-stitched camel leather mojari with ornate brass wire and silk thread embellishments, crafted by heritage cobblers of Jaipur."
    },
    "Maroon Handwoven Cotton Shirt": {
        "name": "Maroon Hand-block Printed Ajrakh Kurta",
        "material": "Handloom Cotton",
        "technique": "Ajrakh block printing",
        "state": "Gujarat",
        "primary_use": "Traditional wear",
        "desc": "Traditional men's kurta crafted from pure handloom cotton and printed with 16-stage ancient Ajrakh resist block printing using natural mineral and vegetable dyes."
    },
    "Brown Handwoven Cotton Shirt": {
        "name": "Brown Hand-block Printed Ajrakh Kurta",
        "material": "Handloom Cotton",
        "technique": "Ajrakh block printing",
        "state": "Gujarat",
        "primary_use": "Traditional wear",
        "desc": "Earthy brown men's ethnic kurta hand-printed with geometric Ajrakh woodblocks in Kutch, combining organic indigo and iron rust dyes."
    },
    "Green Handwoven Cotton Shirt": {
        "name": "Green Hand-block Printed Ajrakh Kurta",
        "material": "Handloom Cotton",
        "technique": "Ajrakh block printing",
        "state": "Gujarat",
        "primary_use": "Traditional wear",
        "desc": "Heritage green Ajrakh handloom cotton kurta featuring traditional star and medallion block prints with mandarin collar styling."
    },
    "Red Handwoven Cotton Shirt": {
        "name": "Red Hand-block Printed Ajrakh Kurta",
        "material": "Handloom Cotton",
        "technique": "Ajrakh block printing",
        "state": "Gujarat",
        "primary_use": "Traditional wear",
        "desc": "Rich crimson red Ajrakh printed cotton kurta hand-dyed with natural madder root using centuries-old Sindhi-Kutchi resist techniques."
    },

    # ── GIFTS CATEGORY ────────────────────────────────────────────
    "Floral Customised Ceramic Mug": {
        "name": "Jaipur Blue Pottery Floral Mug Set",
        "material": "Quartz & Glass",
        "technique": "Blue pottery moulding",
        "state": "Rajasthan",
        "primary_use": "Dining & gifts",
        "desc": "GI-tagged authentic Jaipur blue pottery mugs handcrafted without clay using quartz powder, Fuller's earth, and hand-painted Persian floral motifs in cobalt blue."
    },
    "Festive Customised Ceramic Mug": {
        "name": "Jaipur Blue Pottery Festive Mug Set",
        "material": "Quartz & Glass",
        "technique": "Blue pottery moulding",
        "state": "Rajasthan",
        "primary_use": "Dining & gifts",
        "desc": "Exquisite handcrafted Jaipur blue pottery mug adorned with traditional turquoise and lapis blue glazes, fired at low temperatures for lasting beauty."
    },
    "Yellow Handmade Soy Candle": {
        "name": "Agra Hand-carved Soapstone Jali Aroma Burner",
        "material": "Soapstone",
        "technique": "Stone jali carving",
        "state": "Uttar Pradesh",
        "primary_use": "Home fragrance & decor",
        "desc": "Intricately pierced hand-carved soapstone aroma oil burner with floral lattice patterns, sculpted by master stone artisans in Agra."
    },
    "Blue Handmade Soy Candle": {
        "name": "Bidriware Silver Inlay Keepsake Box",
        "material": "Zinc alloy with pure silver",
        "technique": "Bidriware inlay",
        "state": "Karnataka",
        "primary_use": "Keepsake storage & gifts",
        "desc": "Centuries-old GI craft of Bidar featuring oxidized black metal alloy inlaid with pure 99.9% fine silver wire in intricate Mughal arabesque patterns."
    },
    "Natural Handmade Soy Candle": {
        "name": "Kashmir Walnut Wood Carved Keepsake Box",
        "material": "Seasoned Walnut Wood",
        "technique": "Wood relief carving",
        "state": "Jammu and Kashmir",
        "primary_use": "Jewelry storage & gifts",
        "desc": "Carved from seasoned Kashmiri walnut timber, showcasing 3D relief chinar leaf motifs hand-chiseled with zero nails or modern adhesives."
    },
    "Multicolour Handmade Soy Candle": {
        "name": "Moradabad Engraved Brass Peacock Diya",
        "material": "Solid Brass",
        "technique": "Metal engraving & etching",
        "state": "Uttar Pradesh",
        "primary_use": "Pooja & gifting",
        "desc": "Hand-cast solid brass oil lamp featuring a majestic dancing peacock arch and hand-etched floral engraving from the Brass City of Moradabad."
    },
    "Small Handmade Greeting Card Set": {
        "name": "Madhubani Hand-painted Folk Art Folio",
        "material": "Handmade Cotton Rag Paper",
        "technique": "Madhubani folk painting",
        "state": "Bihar",
        "primary_use": "Gifts & stationery",
        "desc": "Set of archival-grade handmade paper greeting folios hand-painted with bamboo nibs and natural twig brushes using natural Mithila folk pigments."
    },
    "Small Handmade Photo Frame": {
        "name": "Bone & Brass Inlay Handcrafted Picture Frame",
        "material": "Camel Bone & Sheesham Wood",
        "technique": "Bone inlay craft",
        "state": "Rajasthan",
        "primary_use": "Home decor & gifting",
        "desc": "Regal tabletop photo frame crafted by Jodhpur artisans, meticulously embedding hand-cut bone mosaics and brass fillets onto seasoned hardwood."
    },
    "Personalised Keychain": {
        "name": "Dokra Bell Metal Tribal Keychain Charm",
        "material": "Bell Metal / Brass",
        "technique": "Lost-wax casting",
        "state": "Chhattisgarh",
        "primary_use": "Accessories & gifting",
        "desc": "Non-ferrous metal casting using the ancient 4000-year-old Cire-Perdue (lost-wax) technique, showcasing tribal folk motifs from Bastar."
    },
    "Handcrafted Personalised Cushion Cover": {
        "name": "Kutch Rabari Mirror-work Hand-embroidered Cushion Cover",
        "material": "Cotton & Glass Mirrors",
        "technique": "Rabari mirror embroidery",
        "state": "Gujarat",
        "primary_use": "Home decor & gifting",
        "desc": "Heritage square cushion cover adorned with traditional Kutchi Shisha (mirror) work and dense chain stitch embroidery by nomadic Rabari women artisans."
    },

    # ── KIDS CATEGORY ─────────────────────────────────────────────
    "Geometric Handmade Alphabet Blocks": {
        "name": "Channapatna Lacquered Wooden Engine Toy",
        "material": "Ivory Wood (Wrightia Tinctoria)",
        "technique": "Lacquered wood turning",
        "state": "Karnataka",
        "primary_use": "Child play & decor",
        "desc": "100% non-toxic eco-friendly wooden toy train engine turned on lathe and polished using natural vegetable dyes and shellac tree lacquer in Channapatna."
    },
    "Classic Handmade Alphabet Blocks": {
        "name": "Kondapalli Traditional Painted Wooden Toys",
        "material": "Tella Poniki Softwood",
        "technique": "Kondapalli toy carving",
        "state": "Andhra Pradesh",
        "primary_use": "Toys & collectible",
        "desc": "GI-tagged 400-year-old craft from Kondapalli, hand-carved from lightweight softwood and painted with non-toxic earth pigments portraying Indian folklore."
    },
    "Traditional Wooden Puzzle": {
        "name": "Channapatna Handcrafted Wooden 3D Animal Puzzle",
        "material": "Ivory Wood",
        "technique": "Lacquered wood turning",
        "state": "Karnataka",
        "primary_use": "Educational toy",
        "desc": "Handcrafted 3D wooden animal nesting puzzle finished with turmeric and kumkum lacquer polish, perfectly smooth and safe for toddlers."
    },
    "Minimal Wooden Puzzle": {
        "name": "Saharanpur Hand-carved Wooden Tangram Puzzle",
        "material": "Sheesham Rosewood",
        "technique": "Wood cutting & polishing",
        "state": "Uttar Pradesh",
        "primary_use": "Mind puzzle",
        "desc": "Brain-teasing geometric tangram puzzle hand-cut from rich grained sheesham wood with brass inlay accents and smooth hand-sanded edges."
    },
    "Green Cloth Activity Book": {
        "name": "Traditional Kathputli Puppet Doll Set",
        "material": "Cotton Fabric & Wood",
        "technique": "Kathputli puppetry craft",
        "state": "Rajasthan",
        "primary_use": "Play & storytelling",
        "desc": "Handmade Rajasthani string puppet pair carved from mango wood with traditional sequined ghagra choli attire, used in centuries-old folk theater."
    },
    "Red Cloth Activity Book": {
        "name": "Channapatna Wooden Stacking Rainbow Tower",
        "material": "Ivory Wood",
        "technique": "Lacquered wood turning",
        "state": "Karnataka",
        "primary_use": "Toddler motor skills",
        "desc": "Classic 7-tier stacking ring tower hand-turned from Wrightia tinctoria wood, dyed in vibrant natural colors using vegetable extracts."
    },
    "Maroon Cloth Activity Book": {
        "name": "Varanasi Handcrafted Wooden Rolling Bird Toy",
        "material": "Eucalyptus Wood",
        "technique": "Wood turning & lacquering",
        "state": "Uttar Pradesh",
        "primary_use": "Toddler toy",
        "desc": "Traditional hand-turned wooden rolling bird with flapping wings, finished with eco-friendly natural lacquer from the sacred craft clusters of Varanasi."
    },
    "Brown Cloth Activity Book": {
        "name": "Kinhal Handcrafted Wooden Rocking Horse",
        "material": "Kinhal Wood & Tamarind Paste",
        "technique": "Kinhal craft",
        "state": "Karnataka",
        "primary_use": "Toy & collectible",
        "desc": "GI-tagged heritage Kinhal toy horse fashioned with wooden armature and tamarind seed paste relief work, hand-painted in rich traditional tempera colors."
    },
    "Mini Felt Animal Toy": {
        "name": "Rajasthani Handcrafted Stuffed Cotton Elephant",
        "material": "Cotton Fabric & Gotta Patti",
        "technique": "Folk textile stitching",
        "state": "Rajasthan",
        "primary_use": "Toy & decor",
        "desc": "Vibrant handcrafted stuffed cotton elephant embellished with shimmering mirrorwork, colorful thread embroidery, and traditional gotta patti borders."
    },

    # ── HOME & LIVING CATEGORY ────────────────────────────────────
    "Floral Coir Door Mat": {
        "name": "Hand-knotted Kashmir Wool Dhurrie Rug",
        "material": "Pure Indigenous Wool",
        "technique": "Hand knotting",
        "state": "Jammu and Kashmir",
        "primary_use": "Floor decor",
        "desc": "Geometric folk patterned floor rug hand-knotted by generational artisans in Kashmir using hand-spun wool yarn on vertical wooden looms."
    },
    "Festive Coir Door Mat": {
        "name": "Jaipur Hand-block Printed Cotton Dabu Throw",
        "material": "Handloom Cotton",
        "technique": "Dabu mud resist printing",
        "state": "Rajasthan",
        "primary_use": "Living room decor",
        "desc": "Artisan throw blanket crafted using ancient mud-resist Dabu hand block printing and fermented natural indigo dye vats in Bagru."
    },
    "Premium Handwoven Cushion Cover": {
        "name": "Zardozi Embroidered Silk Velvet Cushion Cover",
        "material": "Silk Velvet with Metallic Wire",
        "technique": "Zardozi metal embroidery",
        "state": "Uttar Pradesh",
        "primary_use": "Luxury decor",
        "desc": "Opulent heritage cushion cover adorned with heavy gold and silver metallic wire Zari and French knot work by master Zardozi craftsmen of Lucknow."
    },

    # ── BAGS & ACCESSORIES CATEGORY ───────────────────────────────
    "Classic Leather Handbag": {
        "name": "Shantiniketan Embossed Leather Heritage Handbag",
        "material": "Vegetable-tanned Sheepskin",
        "technique": "Shantiniketan leathercraft",
        "state": "West Bengal",
        "primary_use": "Fashion & travel",
        "desc": "Iconic GI-tagged leather tote featuring traditional batik-inspired embossed scrollwork, crafted by Santiniketan artisan collectives."
    },
    "Geometric Leather Handbag": {
        "name": "Banjara Tribal Embroidered Tote with Mirror Work",
        "material": "Cotton Canvas, Mirrors & Shells",
        "technique": "Banjara needlecraft",
        "state": "Telangana",
        "primary_use": "Fashion bag",
        "desc": "Vibrant boho tote bag created by Lambani/Banjara gypsy artisans using authentic vintage textiles, geometric herringbone embroidery, mirrors, and cowrie shells."
    },
    "Handcrafted Canvas Tote Bag": {
        "name": "Kalamkari Hand-block Printed Canvas Tote Bag",
        "material": "Natural Canvas with Natural Dyes",
        "technique": "Kalamkari block printing",
        "state": "Andhra Pradesh",
        "primary_use": "Daily tote",
        "desc": "Eco-friendly canvas shoulder tote printed with Machilipatnam Kalamkari mythological and floral patterns using alum, myrobalan, and natural madder dyes."
    },
    "Handmade Wallet": {
        "name": "Kutch Hand-embroidered Mirrorwork Zipper Clutch",
        "material": "Mashru Silk & Mirrors",
        "technique": "Kutch hand embroidery",
        "state": "Gujarat",
        "primary_use": "Wallet & clutch",
        "desc": "Spacious zippered clutch wallet embellished with colorful geometric Gujarati embroidery, genuine mirrors, and vibrant handmade pom-poms."
    },

    # ── WOMEN CATEGORY ────────────────────────────────────────────
    "Contemporary Handmade Hair Accessories": {
        "name": "Meenakari Hand-enameled Hair Pin Brooch",
        "material": "Brass with Mineral Enamel",
        "technique": "Meenakari enamel work",
        "state": "Rajasthan",
        "primary_use": "Festive hair styling",
        "desc": "Handcrafted royal hair pin adorned with traditional Jaipur Meenakari enamel art, depicting blooming lotus flowers with pearl droplets."
    },
    "Tribal Handmade Hair Accessories": {
        "name": "Cuttack Tarakasi Silver Filigree Hair Pin",
        "material": "92.5 Sterling Silver Filigree",
        "technique": "Tarakasi wire filigree",
        "state": "Odisha",
        "primary_use": "Bridal & festive wear",
        "desc": "GI-certified Odishan Tarakasi craft featuring wafer-thin sterling silver wires twisted, shaped, and soldered by hand into ethereal hair ornaments."
    }
}

replaced_count = 0
for row_idx in range(2, ws.max_row + 1):
    prod_name = ws.cell(row=row_idx, column=3).value
    if prod_name in REPLACEMENTS:
        rep = REPLACEMENTS[prod_name]
        ws.cell(row=row_idx, column=3).value = rep["name"]
        ws.cell(row=row_idx, column=4).value = rep["material"]
        ws.cell(row=row_idx, column=5).value = rep["technique"]
        ws.cell(row=row_idx, column=6).value = rep["state"]
        ws.cell(row=row_idx, column=11).value = rep["primary_use"]
        ws.cell(row=row_idx, column=13).value = rep["desc"]
        replaced_count += 1

wb.save(str(XLSX_PATH))
print(f"Successfully transformed {replaced_count} rows in {XLSX_PATH}!")
